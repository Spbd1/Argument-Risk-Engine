from __future__ import annotations

from pathlib import Path

from argument_risk_engine.taxonomy.models import RiskSeverity, TaxonomyEntry, TaxonomyPack
from openpyxl import load_workbook

HEADERS = ["id", "name", "description", "severity", "keywords", "examples", "mitigation", "active"]


def import_taxonomy_excel(path: Path | str) -> TaxonomyPack:
    workbook = load_workbook(Path(path))
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(max_row=1))]
    index = {header: idx for idx, header in enumerate(headers)}
    entries: list[TaxonomyEntry] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        data = {header: row[index[header]] if header in index and index[header] < len(row) else None for header in HEADERS}
        entries.append(
            TaxonomyEntry(
                id=str(data["id"] or "").strip(),
                name=str(data["name"] or "").strip(),
                description=str(data["description"] or "").strip(),
                severity=RiskSeverity(str(data.get("severity") or "low").strip().lower()),
                keywords=[part.strip() for part in str(data.get("keywords") or "").split(";") if part.strip()],
                examples=[part.strip() for part in str(data.get("examples") or "").split(";") if part.strip()],
                mitigation=str(data.get("mitigation") or "Escalate for human review."),
                active=str(data.get("active") or "true").strip().lower() not in {"false", "0", "no"},
            )
        )
    return TaxonomyPack(name=Path(path).stem, entries=entries)
